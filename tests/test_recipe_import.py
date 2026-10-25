import io
import json
import re
import tempfile
import unittest
from pathlib import Path
from docx import Document
from werkzeug.datastructures import FileStorage
from recipe_editor import app
from recipe_import import extract_text, parse_recipe_text

TEXT = 'Family cake\nIngredients\n2 cups flour\nDirections\nBake at 350 F.\nServings: 6\nNotes: From Grandma'


class ImportTests(unittest.TestCase):
    def test_sections_and_unclassified_text(self):
        recipe = parse_recipe_text(TEXT)
        self.assertEqual(recipe['ingredients'], ['2 cups flour'])
        self.assertEqual(recipe['instructions'], ['Bake at 350 F.'])
        self.assertEqual(recipe['yield'], '6')
        self.assertEqual(parse_recipe_text('Cake\n2 cups flour')['ingredients'], ['2 cups flour'])

    def test_without_headings_and_wrapped_steps(self):
        recipe = parse_recipe_text("Cake\nA family favorite from 1970\n2 cups flour\n½ cup sugar\nSalt\nVanilla\nMix the ingredients until\nsmooth.\nBake at 350 degrees for 30 minutes.")
        self.assertEqual(recipe['title'], 'Cake')
        self.assertEqual(recipe['ingredients'], ['2 cups flour', '½ cup sugar', 'Salt', 'Vanilla'])
        self.assertEqual(recipe['instructions'], ['Mix the ingredients until smooth.', 'Bake at 350 degrees for 30 minutes.'])
        self.assertEqual(recipe['remarks'], 'A family favorite from 1970')

    def test_missing_directions_heading_and_numbered_steps(self):
        recipe = parse_recipe_text('Bread\nIngredients\n2 cups flour\n1. Mix with water.\n2. Knead well.\nNotes: Add extra water if needed.')
        self.assertEqual(recipe['ingredients'], ['2 cups flour'])
        self.assertEqual(len(recipe['instructions']), 2)
        self.assertEqual(recipe['remarks'], 'Add extra water if needed.')

    def test_explicit_notes_are_not_reclassified(self):
        recipe = parse_recipe_text('Cake\nNotes\n2 cups flour was too much.\nBake longer next time.')
        self.assertEqual(recipe['ingredients'], [])
        self.assertEqual(recipe['instructions'], [])
        self.assertIn('Bake longer next time.', recipe['remarks'])

    def test_word_paragraphs_and_tables(self):
        doc = Document()
        doc.add_paragraph('Cake')
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = 'Ingredients'
        table.cell(1, 0).text = 'Flour'
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        self.assertEqual(extract_text(FileStorage(buffer, filename='recipe.docx')), 'Cake\nIngredients\nFlour')

    def test_input_cleanup_only_after_save(self):
        with tempfile.TemporaryDirectory() as folder:
            content = Path(folder, 'content')
            (content / 'Desserts').mkdir(parents=True)
            inputs = Path(folder, 'inputs')
            inputs.mkdir()
            source = inputs / 'cake.txt'
            source.write_text(TEXT)
            unrelated = inputs / 'other.txt'
            unrelated.write_text('Keep this')
            previous = {key: app.config[key] for key in ('CONTENT_FOLDER', 'INPUT_FOLDER')}
            app.config.update(CONTENT_FOLDER=content, INPUT_FOLDER=inputs)
            try:
                client = app.test_client()
                client.get('/')
                with client.session_transaction() as session:
                    csrf = session['csrf']

                def upload(text):
                    response = client.post('/import', data=dict(csrf=csrf, category='Desserts', document=(io.BytesIO(text.encode()), 'cake.txt')))
                    self.assertEqual(response.status_code, 200)
                    return re.search(r'name="input_token" value="([^"]*)"', response.get_data(as_text=True))[1]

                token = upload(TEXT)
                self.assertTrue(token)
                self.assertTrue(source.exists())
                data = dict(csrf=csrf, category='Desserts', title='Cake', ingredients='', instructions='Bake', input_token=token)
                response = client.post('/submit', data=data)
                self.assertEqual(response.status_code, 400)
                self.assertIn(token, response.get_data(as_text=True))
                self.assertTrue(source.exists())
                data['ingredients'] = 'Flour'
                self.assertEqual(client.post('/submit', data=data).status_code, 302)
                self.assertFalse(source.exists())
                self.assertTrue(unrelated.exists())
                self.assertTrue((content / 'Desserts' / 'Cake.json').exists())

                source.write_text(TEXT)
                self.assertEqual(upload('Different content'), '')
                token = upload(TEXT)
                source.write_text('Replaced source')
                data.update(title='Another cake', input_token=token)
                self.assertEqual(client.post('/submit', data=data).status_code, 302)
                self.assertEqual(source.read_text(), 'Replaced source')
                data.update(title='Third cake', input_token='tampered-token')
                self.assertEqual(client.post('/submit', data=data).status_code, 302)
                self.assertTrue(source.exists())
            finally:
                app.config.update(previous)

    def test_review_save_collision_and_validation(self):
        with tempfile.TemporaryDirectory() as folder:
            Path(folder, 'Desserts').mkdir()
            previous = app.config['CONTENT_FOLDER']
            app.config['CONTENT_FOLDER'] = Path(folder)
            try:
                client = app.test_client()
                self.assertEqual(client.get('/').status_code, 200)
                with client.session_transaction() as session:
                    csrf = session['csrf']
                response = client.post('/import', data=dict(csrf=csrf, category='Desserts', raw=TEXT))
                self.assertEqual(response.status_code, 200)
                self.assertEqual(list(Path(folder).rglob('*.json')), [])
                response = client.post('/import', data=dict(csrf=csrf, category='Desserts', raw='Cake\n2 cups flour\nMix well.\nBake until golden.'))
                html = response.get_data(as_text=True)
                self.assertIn('required>2 cups flour</textarea>', html)
                self.assertIn('required>Mix well.\nBake until golden.</textarea>', html)
                data = dict(csrf=csrf, category='Desserts', title='Cake', ingredients='Flour', instructions='Bake', tags='family, sweet, family')
                self.assertEqual(client.post('/submit', data=data).status_code, 302)
                saved = json.loads(Path(folder, 'Desserts', 'Cake.json').read_text())
                self.assertEqual(saved['tags'], ['family', 'sweet'])
                self.assertEqual(client.post('/submit', data=data).status_code, 409)
                data['title'] = 'Empty'
                data['ingredients'] = ''
                self.assertEqual(client.post('/submit', data=data).status_code, 400)
                self.assertEqual(client.post('/submit', data={}).status_code, 400)
                self.assertEqual(client.get('/?category=../').status_code, 400)
            finally:
                app.config['CONTENT_FOLDER'] = previous


if __name__ == '__main__':
    unittest.main()
